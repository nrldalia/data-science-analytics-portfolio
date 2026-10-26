"""Validated, previewed store imports and bounded store-list queries."""
import copy
import csv
import hashlib
import io
import json
import re
import time
import uuid
import zipfile

from flask import jsonify, request, send_file

from database import begin_write, get_db_connection
from configuration_history import record_version, version_summary
from retailer_archive import archive_state

HEADERS = ('store_id', 'store_name', 'address', 'status', 'exception')
MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 1000
PREVIEW_SECONDS = 30 * 60


def config_hash(config):
    return hashlib.sha256(json.dumps(config, sort_keys=True, separators=(',', ':')).encode()).hexdigest()


def csv_safe(value):
    text = str(value)
    return "'" + text if text.lstrip().startswith(('=', '+', '-', '@')) or text.startswith(('\t', '\r', '\n')) else text


def csv_download(rows, filename):
    stream = io.StringIO(newline='')
    writer = csv.writer(stream)
    for row in rows:
        writer.writerow([csv_safe(v) for v in row])
    return send_file(io.BytesIO(stream.getvalue().encode('utf-8-sig')), mimetype='text/csv',
                     as_attachment=True, download_name=filename)


def parse_file(file):
    if not file or not file.filename:
        raise ValueError('Choose a CSV or Excel (.xlsx) store file.')
    source = file.read(MAX_BYTES + 1)
    if len(source) > MAX_BYTES:
        raise ValueError('Store imports must be at most 2 MB.')
    name = file.filename.lower()
    if name.endswith('.csv'):
        try:
            content = source.decode('utf-8-sig')
            if '\x00' in content:
                raise ValueError('The CSV contains invalid null characters.')
            rows = []
            for values in csv.reader(io.StringIO(content, newline=''), strict=True):
                if len(rows) > MAX_ROWS or len(values) > len(HEADERS):
                    raise ValueError('Use at most 1,000 rows and the five template columns.')
                rows.append(values)
        except (UnicodeDecodeError, csv.Error) as exc:
            raise ValueError('Use a valid UTF-8 CSV file.') from exc
    elif name.endswith('.xlsx'):
        from openpyxl import load_workbook
        try:
            with zipfile.ZipFile(io.BytesIO(source)) as archive:
                if len(archive.infolist()) > 100 or sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValueError('The expanded workbook exceeds the 20 MB limit.')
            book = load_workbook(io.BytesIO(source), read_only=True, data_only=False)
            try:
                sheet = book.worksheets[0]
                if sheet.max_column and sheet.max_column > len(HEADERS) or sheet.max_row and sheet.max_row > MAX_ROWS + 1:
                    raise ValueError('Use at most 1,000 rows and the five template columns; remove formatting outside this range.')
                # Validate actual cell coordinates as dimensions are supplied by the uploader.
                from xml.etree.ElementTree import iterparse
                from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
                with zipfile.ZipFile(io.BytesIO(source)) as archive:
                    with archive.open(sheet._worksheet_path) as document:
                        for _, element in iterparse(document, events=('end',)):
                            if element.tag.rsplit('}', 1)[-1] == 'c':
                                column, row = coordinate_from_string(element.attrib.get('r', ''))
                                if column_index_from_string(column) > len(HEADERS) or row > MAX_ROWS + 1:
                                    raise ValueError('Use at most 1,000 rows and the five template columns.')
                            element.clear()
                sheet.reset_dimensions()
                rows = []
                for cells in sheet.iter_rows(max_row=MAX_ROWS + 1, max_col=len(HEADERS)):
                    if any(c.data_type == 'f' for c in cells):
                        raise ValueError('Formulas are not accepted in store imports; paste values first.')
                    rows.append([c.value for c in cells])
                while rows and not any(v is not None for v in rows[-1]):
                    rows.pop()
                if rows:
                    while rows[0] and rows[0][-1] is None:
                        rows[0].pop()
                    width = len(rows[0])
                    rows[1:] = [values[:width] if not any(v is not None for v in values[width:]) else values for values in rows[1:]]
            finally:
                book.close()
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError('Use a valid, unencrypted Excel (.xlsx) workbook.') from exc
    else:
        raise ValueError('Only CSV and Excel (.xlsx) files are supported.')
    if not rows:
        raise ValueError('The file is empty.')
    header = [str(v or '').strip().lower() for v in rows[0]]
    if len(header) != len(set(header)):
        raise ValueError('Column names must be unique.')
    if not {'store_id', 'store_name'}.issubset(header):
        raise ValueError('Required columns: store_id and store_name. Download the store template.')
    if set(header) - set(HEADERS):
        raise ValueError('Unknown columns. Use store_id, store_name, address, status and exception.')
    records = []
    for number, values in enumerate(rows[1:], 2):
        if not any(v is not None and str(v).strip() for v in values):
            continue
        if len(records) >= MAX_ROWS:
            raise ValueError('Import at most 1,000 stores per retailer.')
        records.append((number, dict(zip(header, values)), len(values) != len(header)))
    if not records:
        raise ValueError('Add at least one store below the header row.')
    return records


def make_preview(config, records, mode):
    existing = {s['key'].casefold(): s for s in config.get('stores', [])}
    counts = {}
    for _, fields, _ in records:
        key = fields.get('store_id')
        if isinstance(key, str):
            folded = key.strip().casefold()
            counts[folded] = counts.get(folded, 0) + 1
    rows, changes = [], []
    summary = dict(total=len(records), added=0, updated=0, unchanged=0, errors=0)
    for number, fields, bad_width in records:
        errors = []
        key, name = fields.get('store_id'), fields.get('store_name')
        if not isinstance(key, str) or not re.fullmatch(r'[A-Za-z0-9_-]{1,64}', key.strip()):
            errors.append('Store ID must contain 1–64 letters, numbers, underscores or hyphens; use text cells.')
        if not isinstance(name, str) or not name.strip() or len(name.strip()) > 120:
            errors.append('Store name must be text containing 1–120 characters.')
        key = key.strip() if isinstance(key, str) else str(key or '')
        name = name.strip() if isinstance(name, str) else str(name or '')
        previous = existing.get(key.casefold())
        if bad_width:
            errors.append('The number of cells does not match the header.')
        if counts.get(key.casefold(), 0) > 1:
            errors.append('Duplicate Store ID within this file (case-insensitive).')
        if previous and mode == 'add':
            errors.append('Store ID already exists. Choose Add and update matching stores to update it.')
        address = fields.get('address', (previous or {}).get('address', ''))
        address = '' if address is None else address
        if not isinstance(address, str) or len(address.strip()) > 2000:
            errors.append('Address must be text of at most 2,000 characters.')
        address = address.strip() if isinstance(address, str) else str(address)
        status = fields.get('status')
        if status is None or isinstance(status, str) and not status.strip():
            status = (previous or {}).get('status', 'Active')
        if not isinstance(status, str) or status.strip().lower() not in ('active', 'inactive'):
            errors.append('Status must be Active or Inactive.')
            status = str(status)
        else:
            status = status.strip().capitalize()
        exception = fields.get('exception')
        if exception is None or isinstance(exception, str) and not exception.strip():
            exception = (previous or {}).get('exception', False)
        elif isinstance(exception, bool):
            pass
        elif isinstance(exception, (int, float)) and exception in (0, 1):
            exception = bool(exception)
        elif isinstance(exception, str) and exception.strip().lower() in ('true', 'false', 'yes', 'no', '1', '0'):
            exception = exception.strip().lower() in ('true', 'yes', '1')
        else:
            errors.append('Exception must be true/false, yes/no or 1/0.')
            exception = str(exception)
        normalized = dict(key=key, name=name, address=address, status=status, exception=exception)
        if previous:
            normalized['id'] = previous['id']
        changed = not previous or any(normalized[k] != previous.get(k, {'address': '', 'status': 'Active', 'exception': False}.get(k)) for k in ('key', 'name', 'address', 'status', 'exception'))
        action = 'error' if errors else ('add' if previous is None else ('update' if changed else 'unchanged'))
        summary[{'error': 'errors', 'add': 'added', 'update': 'updated', 'unchanged': 'unchanged'}[action]] += 1
        rows.append(dict(row=number, action=action, **{k: normalized[k] for k in ('key', 'name', 'address', 'status', 'exception')}, errors=errors))
        if not errors:
            changes.append(normalized)
    if len(existing) + summary['added'] > MAX_ROWS:
        raise ValueError('The resulting retailer would exceed 1,000 stores. No stores were imported.')
    return rows, changes, summary


def register_store_import_routes(app, error):
    with get_db_connection() as db:
        db.execute('CREATE TABLE IF NOT EXISTS store_import_previews(id TEXT PRIMARY KEY,retailer_id TEXT NOT NULL,config_hash TEXT NOT NULL,payload TEXT NOT NULL,expires_at DOUBLE PRECISION NOT NULL,consumed INTEGER NOT NULL DEFAULT 0)')
        db.execute('CREATE INDEX IF NOT EXISTS store_preview_retailer ON store_import_previews(retailer_id,expires_at)')

    @app.get('/api/v1/store-template')
    def store_template():
        kind = request.args.get('format', 'csv')
        if kind == 'csv':
            return csv_download([HEADERS], 'store_import_template.csv')
        if kind == 'xlsx':
            from openpyxl import Workbook
            book = Workbook()
            sheet = book.active
            sheet.title = 'Stores'
            sheet.append(HEADERS)
            sheet.freeze_panes = 'A2'
            sheet.column_dimensions['A'].number_format = '@'
            from openpyxl.worksheet.datavalidation import DataValidation
            for column, choices in [('D', '"Active,Inactive"'), ('E', '"true,false"')]:
                validation = DataValidation(type='list', formula1=choices, allow_blank=True)
                validation.errorTitle = 'Choose a supported value'
                validation.error = 'Use one of the values in the dropdown.'
                validation.showErrorMessage = True
                sheet.add_data_validation(validation)
                validation.add(f'{column}2:{column}1001')
            for col, width in zip('ABCDE', (24, 40, 60, 16, 16)):
                sheet.column_dimensions[col].width = width
            stream = io.BytesIO()
            book.save(stream)
            book.close()
            stream.seek(0)
            return send_file(stream, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name='store_import_template.xlsx')
        return error('Choose csv or xlsx format.')

    @app.get('/api/v1/retailers/<rid>/stores')
    def list_stores(rid):
        try:
            page = int(request.args.get('page', '1'))
            size = int(request.args.get('pageSize', '25'))
            if page < 1 or size not in (10, 25, 50, 100):
                raise ValueError('Use page 1 or higher and pageSize 10, 25, 50 or 100.')
            status = request.args.get('status', 'All')
            sort = request.args.get('sort', 'key')
            direction = request.args.get('direction', 'asc')
            if status not in ('All', 'Active', 'Inactive') or sort not in ('key', 'name', 'suiteCode', 'status') or direction not in ('asc', 'desc'):
                raise ValueError('Choose valid store filters and sorting.')
            query = request.args.get('q', '').strip().casefold()
            if len(query) > 200:
                raise ValueError('Search must be at most 200 characters.')
            with get_db_connection() as db:
                row = db.execute('SELECT payload FROM retailer_configs WHERE id=?', (rid,)).fetchone()
            if not row:
                return error('Retailer not found.', 404)
            stores = json.loads(row['payload']).get('stores', [])
            selected = [s for s in stores if (status == 'All' or s.get('status', 'Active') == status) and (not query or query in ' '.join(str(s.get(k, '')) for k in ('key', 'name', 'suiteCode', 'address')).casefold())]
            selected.sort(key=lambda s: (str(s.get(sort, '')).casefold(), s['key'].casefold()), reverse=direction == 'desc')
            total = len(selected)
            page = min(page, max(1, (total + size - 1) // size))
            return jsonify(success=True, stores=selected[(page - 1) * size:page * size], total=total, totalStores=len(stores), page=page, pageSize=size)
        except (ValueError, TypeError):
            return error('Use valid pagination, filters and sorting (pageSize: 10, 25, 50 or 100).')

    @app.post('/api/v1/retailers/<rid>/stores/import/preview')
    def preview_import(rid):
        try:
            mode = request.form.get('mode', 'add')
            if mode not in ('add', 'upsert'):
                return error('Choose add or upsert import mode.')
            records = parse_file(request.files.get('file'))
            with get_db_connection() as db:
                begin_write(db)
                row = db.execute('SELECT payload FROM retailer_configs WHERE id=?', (rid,)).fetchone()
                if not row:
                    return error('Retailer not found.', 404)
                if archive_state(db, rid)['archived']:
                    return error('Restore this retailer before importing stores.', 409)
                config = json.loads(row['payload'])
                rows, changes, summary = make_preview(config, records, mode)
                stamp = time.time()
                db.execute('DELETE FROM store_import_previews WHERE expires_at<?', (stamp,))
                if db.execute('SELECT COUNT(*) FROM store_import_previews WHERE retailer_id=? AND consumed=0', (rid,)).fetchone()[0] >= 20:
                    return error('Too many pending previews. Confirm one or wait 30 minutes.', 429)
                identity = 'IMPORT-' + uuid.uuid4().hex
                payload = dict(rows=rows, changes=changes, summary=summary, mode=mode)
                db.execute('INSERT INTO store_import_previews VALUES(?,?,?,?,?,0)', (identity, rid, config_hash(config), json.dumps(payload), stamp + PREVIEW_SECONDS))
            return jsonify(success=True, previewId=identity, rows=rows, summary=summary, canConfirm=summary['errors'] == 0 and summary['added'] + summary['updated'] > 0, expiresInSeconds=PREVIEW_SECONDS)
        except ValueError as exc:
            return error(str(exc))

    @app.post('/api/v1/retailers/<rid>/stores/import/confirm')
    def confirm_import(rid):
        from retailer_setup import assign_store_codes, validate_retailer
        body = request.get_json(silent=True)
        if not isinstance(body, dict) or not isinstance(body.get('previewId'), str):
            return error('A previewId is required.')
        with get_db_connection() as db:
            begin_write(db)
            preview = db.execute('SELECT * FROM store_import_previews WHERE id=? AND retailer_id=?', (body['previewId'], rid)).fetchone()
            if not preview:
                return error('Import preview not found. Preview the file again.', 404)
            if preview['consumed']:
                return error('This preview was already imported. Refresh the store list.', 409)
            if preview['expires_at'] <= time.time():
                return error('Import preview expired. Preview the file again.', 410)
            row = db.execute('SELECT payload FROM retailer_configs WHERE id=?', (rid,)).fetchone()
            if not row:
                return error('Retailer not found.', 404)
            if archive_state(db, rid)['archived']:
                return error('Restore this retailer before importing stores.', 409)
            config = json.loads(row['payload'])
            if config_hash(config) != preview['config_hash']:
                return error('Retailer settings changed after this preview. Preview the file again.', 409)
            payload = json.loads(preview['payload'])
            if payload['summary']['errors'] or not payload['summary']['added'] + payload['summary']['updated']:
                return error('Resolve all validation errors and preview changes before importing.')
            updated = copy.deepcopy(config)
            by_key = {s['key'].casefold(): i for i, s in enumerate(updated.get('stores', []))}
            for store in payload['changes']:
                index = by_key.get(store['key'].casefold())
                if index is None:
                    updated['stores'].append(store)
                else:
                    updated['stores'][index] = {**updated['stores'][index], **store}
            try:
                updated = validate_retailer(updated)
                updated['stores'] = assign_store_codes(updated['stores'], config.get('stores', []), db)
            except (ValueError, TypeError) as exc:
                db.rollback()
                return error(str(exc))
            db.execute('UPDATE retailer_configs SET payload=? WHERE id=?', (json.dumps(updated), rid))
            version = record_version(db, rid, updated, 'Stores imported (' + payload['mode'] + ')')
            db.execute('UPDATE store_import_previews SET consumed=1 WHERE id=?', (preview['id'],))
            saved = updated | dict(id=rid, **archive_state(db, rid))
            from processing_readiness import readiness
            saved['processingReadiness'] = readiness(db, rid, saved)
        return jsonify(success=True, retailer=saved, summary=payload['summary'], configurationVersion=version_summary(version))

    @app.get('/api/v1/retailers/<rid>/stores/import/<preview_id>/errors')
    def import_errors(rid, preview_id):
        with get_db_connection() as db:
            preview = db.execute('SELECT payload,expires_at FROM store_import_previews WHERE id=? AND retailer_id=?', (preview_id, rid)).fetchone()
        if not preview:
            return error('Import preview not found.', 404)
        if preview['expires_at'] <= time.time():
            return error('Import preview expired. Preview the file again.', 410)
        rows = [['row', 'store_id', 'store_name', 'errors']]
        rows.extend([r['row'], r['key'], r['name'], '; '.join(r['errors'])] for r in json.loads(preview['payload'])['rows'] if r['errors'])
        return csv_download(rows, 'store_import_errors.csv')
