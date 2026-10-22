"""Input profile validation and bounded parsing, without exposing passwords."""
import csv,io,json,os,re,zipfile,ctypes,fnmatch
from pathlib import Path,PurePosixPath
from hashlib import sha256
DELIMITERS={'Comma':',','Tab':'\t','Semicolon':';','Pipe':'|','Caret':'^'}
ENCODINGS=['utf-8-sig','utf-8','utf-16','utf-16-le','utf-16-be','cp1252','latin-1','ascii']
TYPES=['CSV','TSV','Text','Excel (.xlsx)','JSON']
EXTENSIONS={'CSV':'.csv','TSV':'.tsv','Text':'.txt','Excel (.xlsx)':'.xlsx','JSON':'.json'}
LIMIT=50*1024*1024
def validate_profiles(profiles):
    if not isinstance(profiles,list) or len(profiles)>50:raise ValueError('Use at most 50 input files.')
    from kpi_records import validate_kpi
    result=[];seen=set()
    for p in profiles:
        if not isinstance(p,dict):raise ValueError('Invalid input file settings.')
        pid=p.get('id')
        if not isinstance(pid,str) or not re.fullmatch(r'INPUT-[a-zA-Z0-9-]{1,64}',pid) or pid in seen:raise ValueError('Input files need distinct valid IDs.')
        seen.add(pid)
        for key,values in [('delimiter',DELIMITERS),('encoding',ENCODINGS),('fileType',TYPES),('purpose',['Transaction','Master Product']),('readMode',['Manual','Automatic'])]:
            if p.get(key) not in values:raise ValueError(f'Choose a valid {key}.')
        modern=p.get('settingsVersion')==2
        for key in ['filename','filePath']:
            if not isinstance(p.get(key),str) or (not p[key].strip() and (key=='filename' or modern or p.get('readMode')=='Automatic')) or len(p[key])>(200 if key=='filename' else 1000):raise ValueError(f'{key} is required and must fit the length limit.')
        if any(c in p['filename'] for c in ['/', '\\', ':']) or p['filename'] in ['.','..']:raise ValueError('Filename must be a name or wildcard pattern, without a path.')
        if not isinstance(p.get('isZipFolder'),bool):raise ValueError('Choose whether the input is a ZIP archive.')
        for key,maximum in [('maxColumns',1000),('headerRow',10000)]:
            if isinstance(p.get(key),bool) or not isinstance(p.get(key),int) or not 1<=p[key]<=maximum:raise ValueError(f'{key} must be a whole number from 1 to {maximum}.')
        if p.get('watchInterval',60) not in [60,300,900]:raise ValueError('Choose a watch interval of 1, 5 or 15 minutes.')
        zipname=p.get('zipFolderName','')
        if not isinstance(zipname,str) or len(zipname)>200 or any(c in zipname for c in ['/', '\\', ':']):raise ValueError('Zip Folder Name must be a filename or pattern, without a path.')
        if modern and p['isZipFolder'] and (not zipname.strip() or not zipname.lower().endswith('.zip')):raise ValueError('Zip Folder Name is required and must end in .zip.')
        revision=p.get('passwordRevision',0)
        if isinstance(revision,bool) or not isinstance(revision,int) or not 0<=revision<=2147483647:raise ValueError('Invalid password revision.')
        if p['fileType']=='JSON' and p['headerRow']!=1:raise ValueError('JSON has named keys; Header Row must be 1.')
        result.append({k:p[k] for k in ['id','delimiter','encoding','fileType','purpose','filename','filePath','isZipFolder','maxColumns','headerRow','readMode']} | {'watchInterval':p.get('watchInterval',60),'passwordRevision':int(p.get('passwordRevision',0)), 'kpi':validate_kpi(p.get('kpi')), 'settingsVersion':p.get('settingsVersion',1), 'zipFolderName':zipname, 'hasSheetPassword':bool(p.get('hasSheetPassword',False))})
    return result

def parse_hash(profile):
    values={k:profile.get(k) for k in ['delimiter','encoding','fileType','isZipFolder','maxColumns','headerRow','passwordRevision']}
    if profile.get('settingsVersion')==2 and profile['isZipFolder']:values['innerFilename']=profile['filename']
    return sha256(json.dumps(values,sort_keys=True).encode()).hexdigest()

def protect_password(text):
    if os.name!='nt':raise ValueError('Password storage requires Windows user encryption.')
    from ctypes import wintypes
    class Blob(ctypes.Structure):_fields_=[('cbData',wintypes.DWORD),('pbData',ctypes.POINTER(ctypes.c_ubyte))]
    source=text.encode();buffer=ctypes.create_string_buffer(source);a=Blob(len(source),ctypes.cast(buffer,ctypes.POINTER(ctypes.c_ubyte)));b=Blob()
    if not ctypes.windll.crypt32.CryptProtectData(ctypes.byref(a),None,None,None,None,1,ctypes.byref(b)):raise ValueError('Unable to securely save the sheet password.')
    try:return ctypes.string_at(b.pbData,b.cbData)
    finally:ctypes.windll.kernel32.LocalFree(b.pbData)

def reveal_password(protected):
    if not protected:return ''
    from ctypes import wintypes
    class Blob(ctypes.Structure):_fields_=[('cbData',wintypes.DWORD),('pbData',ctypes.POINTER(ctypes.c_ubyte))]
    buffer=ctypes.create_string_buffer(protected);a=Blob(len(protected),ctypes.cast(buffer,ctypes.POINTER(ctypes.c_ubyte)));b=Blob()
    if not ctypes.windll.crypt32.CryptUnprotectData(ctypes.byref(a),None,None,None,None,1,ctypes.byref(b)):raise ValueError('Unable to decrypt the saved sheet password under this Windows user.')
    try:return ctypes.string_at(b.pbData,b.cbData).decode()
    finally:ctypes.windll.kernel32.LocalFree(b.pbData)

def read_records(source,filename,p,password='',limit=LIMIT):
    if len(source)>limit:raise ValueError('Input exceeds the file size limit.')
    extension=EXTENSIONS[p['fileType']]
    if p['isZipFolder']:
        try:
            with zipfile.ZipFile(io.BytesIO(source)) as z:
                members=z.infolist()
                if len(members)>100:raise ValueError('ZIP supports at most 100 entries.')
                for member in members:
                    name=member.filename.replace('\\','/');parts=PurePosixPath(name)
                    if parts.is_absolute() or '..' in parts.parts or ':' in name:raise ValueError('ZIP contains an unsafe entry name.')
                selected=[m for m in members if not m.is_dir() and m.filename.lower().endswith(extension) and (p.get('settingsVersion')!=2 or fnmatch.fnmatchcase(PurePosixPath(m.filename.replace('\\','/')).name,p['filename']))]
                if len(selected)!=1:raise ValueError('ZIP must contain exactly one matching Filename of the configured type.')
                member=selected[0]
                if member.flag_bits&1:raise ValueError('Password-protected ZIP archives are not supported; Sheet Password is for Excel.')
                if member.file_size>limit:raise ValueError('Uncompressed ZIP input exceeds the file size limit.')
                with z.open(member) as f:source=f.read(limit+1)
                if len(source)>limit:raise ValueError('Uncompressed ZIP input exceeds the file size limit.')
        except zipfile.BadZipFile:raise ValueError('Input is not a valid ZIP archive.')
        except (NotImplementedError,RuntimeError):raise ValueError('ZIP compression or protection is unsupported.')
    elif not filename.lower().endswith(extension):raise ValueError('Sample or input file must match the configured file type.')
    maxcols=p['maxColumns'];header=p['headerRow']
    if p['fileType']=='JSON':
        try:records=json.loads(source.decode(p['encoding']))
        except (UnicodeError,json.JSONDecodeError):raise ValueError('JSON cannot be read with the chosen encoding or structure.')
        records=[records] if isinstance(records,dict) else records
    else:
        if p['fileType']=='Excel (.xlsx)':
            import msoffcrypto
            try:
                office=msoffcrypto.OfficeFile(io.BytesIO(source))
                if office.is_encrypted():
                    if not password:raise ValueError('This Excel workbook requires a Sheet Password.')
                    output=io.BytesIO();office.load_key(password=password);office.decrypt(output);source=output.getvalue()
            except ValueError:raise
            except Exception:raise ValueError('Excel could not be opened. Check the workbook and Sheet Password.')
            from openpyxl import load_workbook
            try:workbook=load_workbook(io.BytesIO(source),read_only=True,data_only=True)
            except Exception:raise ValueError('Excel workbook could not be read.')
            try:
                sheet=workbook.worksheets[0]
                if sheet.max_column and sheet.max_column>maxcols:raise ValueError('Excel input exceeds Max Columns.')
                rows=sheet.iter_rows(values_only=True)
                records=tabular_records(rows,header,maxcols)
            finally:workbook.close()
        else:
            try:text=source.decode(p['encoding'])
            except UnicodeError:raise ValueError('Input cannot be decoded with the chosen encoding.')
            try:records=tabular_records(csv.reader(io.StringIO(text,newline=''),delimiter=DELIMITERS[p['delimiter']],strict=True),header,maxcols)
            except csv.Error:raise ValueError('Delimited input is malformed; check delimiter and quoting.')
    if not isinstance(records,list) or not records or not all(isinstance(r,dict) for r in records):raise ValueError('Input must contain at least one data row.')
    columns=set(k for r in records for k in r)
    if any(not isinstance(k,str) or not k.strip() for k in columns):raise ValueError('Headers or JSON keys must be nonempty text.')
    if len(columns)>maxcols:raise ValueError('Input exceeds Max Columns.')
    return records

def tabular_records(rows,header,maxcols):
    iterator=iter(rows)
    for _ in range(header-1):
        if next(iterator,None) is None:raise ValueError('Header Row is beyond the end of the input.')
    raw=next(iterator,None)
    if raw is None:raise ValueError('Header Row is beyond the end of the input.')
    headers=[str(v).strip() if v is not None else '' for v in raw]
    if not headers or len(headers)>maxcols:raise ValueError('Header exceeds Max Columns or is empty.')
    if any(not h for h in headers) or len(set(headers))!=len(headers):raise ValueError('Headers must be nonempty and unique.')
    records=[]
    for row in iterator:
        if not any(v is not None and str(v)!='' for v in row):continue
        if len(row)>maxcols or len(row)>len(headers):raise ValueError('A row has more columns than the header or Max Columns.')
        values=[v.isoformat() if hasattr(v,'isoformat') else str(v) if v is not None else '' for v in row]
        records.append(dict(zip(headers,values+['']*(len(headers)-len(values)))))
    return records
