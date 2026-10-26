import csv
import io
import json
import unittest
from unittest.mock import patch

import test_store_management as fixtures
from database import get_db_connection


class StoreImportTests(unittest.TestCase):
    setUp = fixtures.StoreManagementTests.setUp

    def preview(self, text, mode='add', rid=None, name='stores.csv'):
        return self.client.post(f'/api/v1/retailers/{rid or self.rid}/stores/import/preview',
                                data={'file': (io.BytesIO(text.encode() if isinstance(text, str) else text), name), 'mode': mode})

    def confirm(self, preview):
        return self.client.post(f'/api/v1/retailers/{self.rid}/stores/import/confirm', json={'previewId': preview['previewId']})

    def test_preview_then_confirm_preserves_existing_and_history(self):
        before = self.saved['stores'][0]
        preview = self.preview('store_id,store_name,address,status,exception\nS2,Second,2 Demo Road,Active,false\n').get_json()
        self.assertEqual(preview['summary'], dict(total=1, added=1, updated=0, unchanged=0, errors=0))
        self.assertEqual(self.client.get(f'/api/v1/retailers/{self.rid}/stores').json['total'], 1)
        result = self.confirm(preview)
        self.assertEqual(result.status_code, 200, result.data)
        stores = result.json['retailer']['stores']
        self.assertEqual(stores[0], before)
        self.assertNotEqual(stores[1]['suiteCode'], before['suiteCode'])
        self.assertEqual(result.json['configurationVersion']['reason'], 'Stores imported (add)')
        self.assertEqual(self.confirm(preview).status_code, 409)

    def test_upsert_preserves_identity_and_missing_optional_fields(self):
        old = self.saved['stores'][0]
        preview = self.preview('store_id,store_name\ns1,Updated Name\n', 'upsert').json
        self.assertEqual(preview['summary']['updated'], 1)
        result = self.confirm(preview)
        self.assertEqual(result.status_code, 200, result.data)
        store = result.json['retailer']['stores'][0]
        self.assertEqual(store['id'], old['id'])
        self.assertEqual(store['suiteCode'], old['suiteCode'])
        self.assertEqual(store['address'], old['address'])
        self.assertEqual(store['status'], old['status'])
        self.assertEqual(store['name'], 'Updated Name')

    def test_invalid_rows_block_all_changes_and_error_export_is_safe(self):
        preview = self.preview('store_id,store_name,status,exception\nS2,=HYPERLINK(123),Unknown,maybe\ns2,Duplicate,Active,false\n').json
        self.assertFalse(preview['canConfirm'])
        self.assertEqual(preview['summary']['errors'], 2)
        self.assertEqual(self.confirm(preview).status_code, 400)
        report = self.client.get(f"/api/v1/retailers/{self.rid}/stores/import/{preview['previewId']}/errors")
        self.assertEqual(report.status_code, 200)
        rows = list(csv.reader(io.StringIO(report.data.decode('utf-8-sig'))))
        self.assertEqual(rows[1][2], "'=HYPERLINK(123)")
        self.assertEqual(self.client.get(f'/api/v1/retailers/{self.rid}/stores').json['total'], 1)

    def test_stale_expired_and_wrong_retailer_previews(self):
        preview = self.preview('store_id,store_name\nS2,Second\n').json
        wrong = self.client.post('/api/v1/retailers/not-this-retailer/stores/import/confirm', json={'previewId': preview['previewId']})
        self.assertEqual(wrong.status_code, 404)
        self.assertEqual(self.client.put(f'/api/v1/retailers/{self.rid}', json={**self.saved, 'name': 'Changed'}).status_code, 200)
        self.assertEqual(self.confirm(preview).status_code, 409)
        preview = self.preview('store_id,store_name\nS2,Second\n').json
        with get_db_connection() as db:
            db.execute('UPDATE store_import_previews SET expires_at=0 WHERE id=?', (preview['previewId'],))
        self.assertEqual(self.confirm(preview).status_code, 410)

    def test_xlsx_template_and_formula_rejection(self):
        from openpyxl import load_workbook
        template = self.client.get('/api/v1/store-template?format=xlsx')
        book = load_workbook(io.BytesIO(template.data))
        book.active.append(['S2', 'Second', '', 'Active', False])
        stream = io.BytesIO();book.save(stream)
        preview = self.preview(stream.getvalue(), name='stores.xlsx')
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertTrue(preview.json['canConfirm'])
        book.active['B2'] = '=1+1'
        stream = io.BytesIO();book.save(stream);book.close()
        self.assertEqual(self.preview(stream.getvalue(), name='stores.xlsx').status_code, 400)

    def test_xlsx_spoofed_dimensions_and_typed_invalid_values(self):
        from datetime import datetime
        from openpyxl import Workbook
        import re
        import zipfile
        book = Workbook()
        book.active.append(['store_id', 'store_name', 'status', 'exception'])
        book.active.append(['S2', 'Second', datetime(2026, 1, 1), datetime(2026, 1, 2)])
        stream = io.BytesIO();book.save(stream)
        result = self.preview(stream.getvalue(), name='stores.xlsx')
        self.assertEqual(result.status_code, 200, result.data)
        self.assertEqual(result.json['summary']['errors'], 1)
        book.active['A1002'] = 'HIDDEN-OVERFLOW'
        stream = io.BytesIO();book.save(stream);book.close()
        manipulated = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(stream.getvalue())) as original, zipfile.ZipFile(manipulated, 'w') as target:
            for item in original.infolist():
                content = original.read(item.filename)
                if item.filename == 'xl/worksheets/sheet1.xml':
                    content = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:D2"', content)
                target.writestr(item, content)
        self.assertEqual(self.preview(manipulated.getvalue(), name='stores.xlsx').status_code, 400)

    def test_format_limits_and_unknown_headers(self):
        self.assertEqual(self.preview('store_id,store_name,unexpected\nS2,Second,yes\n').status_code, 400)
        self.assertEqual(self.preview('store_id,store_id,store_name\nS2,S3,Second\n').status_code, 400)
        self.assertEqual(self.preview('store_id,store_name\n').status_code, 400)
        self.assertEqual(self.preview('a' * (2 * 1024 * 1024 + 1)).status_code, 400)
        self.assertEqual(self.preview('store_id,store_name\n' + ''.join(f'NEW-{i},Name\n' for i in range(1001))).status_code, 400)

    def test_sequence_exhaustion_rolls_back_every_allocation(self):
        preview = self.preview('store_id,store_name\nNEW1,New one\nNEW2,New two\n').json
        with get_db_connection() as db:
            before = db.execute('SELECT COUNT(*) FROM store_code_registry').fetchone()[0]
            original = db.execute('SELECT last_value FROM store_code_sequence WHERE id=1').fetchone()[0]
            db.execute('UPDATE store_code_sequence SET last_value=9999998 WHERE id=1')
        try:
            result = self.confirm(preview)
            self.assertEqual(result.status_code, 400, result.data)
            with get_db_connection() as db:
                self.assertEqual(db.execute('SELECT COUNT(*) FROM store_code_registry').fetchone()[0], before)
                self.assertEqual(db.execute('SELECT last_value FROM store_code_sequence WHERE id=1').fetchone()[0], 9999998)
        finally:
            with get_db_connection() as db:
                db.execute('UPDATE store_code_sequence SET last_value=? WHERE id=1', (original,))

    def test_no_change_preview_and_archival_have_no_save(self):
        preview = self.preview('store_id,store_name\nS1,Main\n', 'upsert').json
        self.assertEqual(preview['summary']['unchanged'], 1)
        self.assertFalse(preview['canConfirm'])
        self.assertEqual(self.confirm(preview).status_code, 400)
        self.assertEqual(self.client.post('/api/v1/retailers/archive', json={'ids': [self.rid]}).status_code, 200)
        self.assertEqual(self.preview('store_id,store_name\nS2,Second\n').status_code, 409)

    def test_paginated_search_filter_sort_and_bounds(self):
        text = 'store_id,store_name,address,status\n' + ''.join(f'S{i:03d},Branch {i:03d},Town {i % 3},{"Active" if i % 2 else "Inactive"}\n' for i in range(2, 302))
        response = self.confirm(self.preview(text).json)
        self.assertEqual(response.status_code, 200, response.data)
        page = self.client.get(f'/api/v1/retailers/{self.rid}/stores?page=2&pageSize=25&sort=name&direction=desc').json
        self.assertEqual(page['total'], 301);self.assertEqual(len(page['stores']), 25)
        filtered = self.client.get(f'/api/v1/retailers/{self.rid}/stores?q=Town%201&status=Active&pageSize=100').json
        self.assertTrue(all('Town 1' in s['address'] and s['status'] == 'Active' for s in filtered['stores']))
        self.assertEqual(self.client.get(f'/api/v1/retailers/{self.rid}/stores?pageSize=999').status_code, 400)
        self.assertEqual(self.client.get(f'/api/v1/retailers/{self.rid}/stores?sort=payload').status_code, 400)


if __name__ == '__main__':
    unittest.main()
